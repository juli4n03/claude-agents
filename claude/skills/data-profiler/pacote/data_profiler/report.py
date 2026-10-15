"""Exportação dos resultados (Excel, Parquet, HTML/gráficos). Dependências opcionais degradam com warning."""

from __future__ import annotations

import json
import os
from typing import TYPE_CHECKING, Any, Sequence

import polars as pl

from ._utils import garantir_dependencia, logger

if TYPE_CHECKING:  # pragma: no cover
    from .correlation import CorrelacaoResult
    from .predictive_power import PoderPreditivoResult
    from .profiling import ProfileResult

_VERMELHO = "FFC7CE"
_AMARELO = "FFEB9C"
_VERDE = "C6EFCE"


def _escrever_abas(path: str, abas: dict[str, pl.DataFrame | None], destaques: dict[str, dict[str, dict[str, str]]]) -> str:
    """Escreve várias abas com ``openpyxl`` e aplica formatação condicional por valor exato.

    ``destaques = {aba: {coluna: {valor: cor_hex}}}``.
    """
    openpyxl = garantir_dependencia("openpyxl", "exportacao para Excel")
    if openpyxl is None:
        return ""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, df in abas.items():
        if df is None:
            continue
        ws = wb.create_sheet(nome[:31])
        cols = df.columns
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in df.iter_rows():
            ws.append([_celula(v) for v in row])
        regras = destaques.get(nome, {})
        for col, mapa in regras.items():
            if col not in cols:
                continue
            j = cols.index(col) + 1
            for i in range(2, ws.max_row + 1):
                cel = ws.cell(row=i, column=j)
                cor = mapa.get(str(cel.value))
                if cor:
                    cel.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
        for j, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(12, len(col) + 2), 45)
        ws.freeze_panes = "B2"
    wb.save(path)
    logger.info("Excel salvo em %s", path)
    return path


def _celula(v: Any) -> Any:
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, float) and (v != v):
        return None
    return v


def _metadados_df(meta: dict) -> pl.DataFrame:
    return pl.DataFrame({"chave": list(meta), "valor": [json.dumps(v, ensure_ascii=False, default=str) if not isinstance(v, (str, int, float, bool)) or v is None else str(v) for v in meta.values()]})


# ----------------------------------------------------------------------------- profiling
def profile_para_excel(res: "ProfileResult", path: str) -> str:
    eleita = {"Nao Eleita": _VERMELHO, "Eleita": _VERDE}
    abas = {"resumo": res.resumo(), "numericas": res.numericas, "categoricas": res.categoricas, "datas_resumo": res.datas_resumo, "datas_safra": res.datas_safra, "estabilidade": res.estabilidade, "psi": res.psi, "bad_rate_safra": res.bad_rate_safra, "metadados": _metadados_df(res.metadados)}
    destaques = {"numericas": {"variavel_eleita": eleita}, "categoricas": {"variavel_eleita": eleita}, "resumo": {"variavel_eleita": eleita}, "psi": {"classificacao_psi": {"instavel": _VERMELHO, "atencao": _AMARELO, "estavel": _VERDE}}}
    return _escrever_abas(path, abas, destaques)


def profile_para_parquet(res: "ProfileResult", diretorio: str) -> list[str]:
    os.makedirs(diretorio, exist_ok=True)
    saida = []
    for nome, df in (("numericas", res.numericas), ("categoricas", res.categoricas), ("datas_resumo", res.datas_resumo), ("datas_safra", res.datas_safra), ("estabilidade", res.estabilidade), ("psi", res.psi), ("bad_rate_safra", res.bad_rate_safra)):
        if df is not None:
            p = os.path.join(diretorio, f"{nome}.parquet")
            df.write_parquet(p)
            saida.append(p)
    with open(os.path.join(diretorio, "metadados.json"), "w", encoding="utf-8") as f:
        json.dump(res.metadados, f, ensure_ascii=False, indent=2, default=str)
    return saida


# ----------------------------------------------------------------------------- poder preditivo
def poder_preditivo_para_excel(res: "PoderPreditivoResult", path: str) -> str:
    rec = {"Selecionar": _VERDE, "Revisar": _AMARELO, "Descartar": _VERMELHO}
    abas = {"resumo": res.resumo, "faixas": res.faixas, "por_safra": res.por_safra, "importancia": res.importancia, "cortes": pl.DataFrame({"variavel": list(res.cortes), "cortes_json": [json.dumps(v, ensure_ascii=False, default=str) for v in res.cortes.values()]}), "metadados": _metadados_df(res.metadados)}
    return _escrever_abas(path, abas, {"resumo": {"recomendacao": rec, "classificacao_iv": {"suspeito": _VERMELHO, "forte": _VERDE}}})


def grafico_variavel(res: "PoderPreditivoResult", variavel: str):
    """Barras de volume por faixa + linha de bad rate; e IV/Gini por safra (se houver). Devolve figuras plotly."""
    go = garantir_dependencia("plotly.graph_objects", "graficos")
    if go is None:
        return None
    from plotly.subplots import make_subplots

    f = res.faixas.filter(pl.col("variavel") == variavel).sort("ordem_faixa")
    if not f.height:
        raise KeyError(f"Variável '{variavel}' não encontrada nas faixas.")
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    fig.add_trace(go.Bar(x=f["faixa"].to_list(), y=f["pct_qtd"].to_list(), name="% volume"), secondary_y=False)
    if "bad_rate" in f.columns:
        fig.add_trace(go.Scatter(x=f["faixa"].to_list(), y=f["bad_rate"].to_list(), name="bad rate (%)", mode="lines+markers"), secondary_y=True)
    elif "media_target" in f.columns:
        fig.add_trace(go.Scatter(x=f["faixa"].to_list(), y=f["media_target"].to_list(), name="média do target", mode="lines+markers"), secondary_y=True)
    fig.update_layout(title=f"{variavel}: volume por faixa e target", xaxis_title="faixa", template="plotly_white")
    figs = [fig]
    if res.por_safra is not None and res.por_safra.height:
        s = res.por_safra.filter(pl.col("variavel") == variavel).sort("safra")
        if s.height:
            fig2 = go.Figure()
            fig2.add_trace(go.Scatter(x=s["safra"].to_list(), y=s["iv"].to_list(), name="IV", mode="lines+markers"))
            fig2.add_trace(go.Scatter(x=s["safra"].to_list(), y=s["gini"].to_list(), name="Gini", mode="lines+markers"))
            fig2.update_layout(title=f"{variavel}: IV e Gini por safra", template="plotly_white")
            figs.append(fig2)
    return figs


def poder_preditivo_para_html(res: "PoderPreditivoResult", path: str, variaveis: Sequence[str] | None = None) -> str | None:
    if garantir_dependencia("plotly", "relatorio HTML") is None:
        return None
    vars_ = list(variaveis or res.resumo["variavel"].to_list())
    partes = ["<html><head><meta charset='utf-8'><title>Poder preditivo</title></head><body>", "<h1>Poder preditivo</h1>"]
    partes.append(res.resumo.select([c for c in ("ranking", "variavel", "recomendacao", "iv", "classificacao_iv", "ks", "gini_woe", "obs") if c in res.resumo.columns]).to_pandas().to_html(index=False))
    primeiro = True
    for v in vars_:
        figs = grafico_variavel(res, v) or []
        for fig in figs:
            partes.append(fig.to_html(full_html=False, include_plotlyjs="cdn" if primeiro else False))
            primeiro = False
    partes.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(partes))
    return path


# ----------------------------------------------------------------------------- correlação
def correlacao_para_excel(res: "CorrelacaoResult", path: str) -> str:
    abas: dict[str, pl.DataFrame | None] = {"pares": res.pares, "pares_altos": res.pares_altos, "vif": res.vif}
    for m, mat in res.matrizes.items():
        abas[f"matriz_{m}"] = mat
    abas["metadados"] = _metadados_df(res.metadados)
    return _escrever_abas(path, abas, {"vif": {"multicolinearidade": {"alta": _VERMELHO, "moderada": _AMARELO}}})
